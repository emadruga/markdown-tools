#!/usr/bin/env python3
"""
Markdown → DOCX Beautifier

Converts markdown lab files to beautifully formatted DOCX documents,
matching the visual style of Lab1_e_Lab2.docx.

Usage:
    python markdown2docx_beautifier.py <input.md>

The DOCX is saved in the same directory as the input file.
"""

import argparse
import logging
import os
import re
import sys
from pathlib import Path

_level = getattr(logging, os.environ.get('LOGLEVEL', 'INFO').upper(), logging.INFO)
logging.basicConfig(format='%(levelname)s: %(message)s', level=_level)
log = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
except ImportError:
    print("Error: python-docx is not installed. Install with: pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ──────────────────────────────────────────────
# Color palette
# ──────────────────────────────────────────────
C_PRIMARY      = '1F4E79'
C_HEADING3     = '1F4D78'
C_ACCENT       = '2E74B5'
C_BODY         = '000000'
C_NOTE         = '444444'
C_MUTED        = '888888'
C_PLACEHOLDER  = 'BBBBBB'
C_SUBTITLE     = '595959'
C_WHITE        = 'FFFFFF'

BG_HEADER_DARK   = '1F4E79'
BG_HEADER_MEDIUM = '2E74B5'
BG_HEADER_LIGHT  = 'DAEAF7'
BG_CODE          = 'F4F4F4'
BG_TABLE_ALT     = 'F9F9F9'
BG_INFO          = 'E3F2FD'
BG_TIP           = 'E0F2F1'
BG_WARNING       = 'FFF2CC'
BG_DANGER        = 'FCE4E4'
BG_ANSWER        = 'FAFAFA'
BG_PLACEHOLDER   = 'F0F0F0'
BG_CHECKID       = '1F3864'
BG_CHECKDESC     = 'EBF3FB'

BADGE_COLORS = {
    'Iniciante':     ('1E7A1E', C_WHITE),
    'Intermediário': ('C45911', C_WHITE),
    'Avançado':      ('C00000', C_WHITE),
}

BORDER_CHAPTER = '1F3864'
BORDER_SECTION = '2E74B5'
BORDER_TIP     = '00695C'


# ──────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────
_bookmark_counter = 99  # will start at 100

def next_bookmark_id():
    global _bookmark_counter
    _bookmark_counter += 1
    return _bookmark_counter


def hex_to_rgb(h):
    h = h.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_cell_bg(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_paragraph_bg(paragraph, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}" w:val="clear"/>')
    paragraph._p.get_or_add_pPr().append(shading)


def set_paragraph_border_left(paragraph, color, sz='6'):
    pPr = paragraph._p.get_or_add_pPr()
    borders = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="{sz}" w:color="{color}" w:space="4"/>'
        f'</w:pBdr>'
    )
    pPr.append(borders)


def set_paragraph_border_bottom(paragraph, color, sz='6'):
    pPr = paragraph._p.get_or_add_pPr()
    borders = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:color="{color}" w:space="1"/>'
        f'</w:pBdr>'
    )
    pPr.append(borders)


def add_table_borders(table, color='auto', sz='4'):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:color="{color}" w:space="0"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:color="{color}" w:space="0"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:color="{color}" w:space="0"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:color="{color}" w:space="0"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:color="{color}" w:space="0"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:color="{color}" w:space="0"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))


def add_run(paragraph, text, bold=False, italic=False, size=None, color=None, font_name=None):
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = hex_to_rgb(color)
    if font_name:
        run.font.name = font_name
    return run


def add_bookmark(paragraph, name, bookmark_id):
    tag_start = parse_xml(
        f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{name}"/>'
    )
    tag_end = parse_xml(
        f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>'
    )
    paragraph._p.append(tag_start)
    paragraph._p.append(tag_end)


def add_hyperlink(paragraph, anchor, text, color='0563C1'):
    """Insert an internal hyperlink (w:hyperlink with w:anchor)."""
    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} w:anchor="{anchor}"/>'
    )
    run_xml = parse_xml(
        f'<w:r {nsdecls("w")}>'
        f'  <w:rPr>'
        f'    <w:color w:val="{color}"/>'
        f'    <w:u w:val="single"/>'
        f'    <w:sz w:val="{11 * 2}"/>'
        f'  </w:rPr>'
        f'  <w:t xml:space="preserve">{_xml_escape(text)}</w:t>'
        f'</w:r>'
    )
    hyperlink.append(run_xml)
    paragraph._p.append(hyperlink)


def _xml_escape(text):
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


def _sanitize_bookmark(name):
    """Sanitize a bookmark name for Word compatibility.

    Word bookmark rules:
    - Only ASCII letters, digits, and underscores
    - Must start with a letter
    - Max 40 characters
    """
    import unicodedata
    # Normalize unicode (e.g. á → a)
    nfkd = unicodedata.normalize('NFKD', name)
    ascii_name = ''.join(c for c in nfkd if not unicodedata.combining(c))
    # Replace hyphens and spaces with underscores
    ascii_name = re.sub(r'[-\s]+', '_', ascii_name)
    # Remove anything that isn't alphanumeric or underscore
    ascii_name = re.sub(r'[^A-Za-z0-9_]', '', ascii_name)
    # Must start with a letter
    if ascii_name and not ascii_name[0].isalpha():
        ascii_name = 'bk_' + ascii_name
    # Truncate to 40 chars
    ascii_name = ascii_name[:40]
    return ascii_name


def add_badge(paragraph, level):
    bg, fg = BADGE_COLORS.get(level, (BG_HEADER_MEDIUM, C_WHITE))
    run = paragraph.add_run(f'  {level}  ')
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = hex_to_rgb(fg)
    run.font.name = 'Arial'
    # shading on run
    rPr = run._r.get_or_add_rPr()
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}" w:val="clear"/>')
    rPr.append(shading)


def set_cell_width(cell, width_cm):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(width_cm * 567)}" w:type="dxa"/>')
    tcPr.append(tcW)


def set_cell_vertical_alignment(cell, val='center'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{val}"/>')
    tcPr.append(vAlign)


# ──────────────────────────────────────────────
# Markdown Parser
# ──────────────────────────────────────────────

class MarkdownParser:
    """Parse markdown into a list of blocks."""

    def __init__(self, text):
        self.lines = text.split('\n')
        self.pos = 0
        self.blocks = []

    def parse(self):
        while self.pos < len(self.lines):
            line = self.lines[self.pos]

            # Blank line
            if not line.strip():
                self.pos += 1
                continue

            # Horizontal rule
            if re.match(r'^---+\s*$', line.strip()):
                self.blocks.append({'type': 'hr'})
                self.pos += 1
                continue

            # Code block
            if line.strip().startswith('```'):
                self._parse_code_block()
                continue

            # Heading
            m = re.match(r'^(#{1,6})\s+(.*)', line)
            if m:
                level = len(m.group(1))
                text = m.group(2).strip()
                self.blocks.append({'type': 'heading', 'level': level, 'text': text})
                self.pos += 1
                continue

            # Blockquote
            if line.strip().startswith('>'):
                self._parse_blockquote()
                continue

            # Table
            if '|' in line and self.pos + 1 < len(self.lines) and re.search(r'\|[\s:]*-+', self.lines[self.pos + 1]):
                self._parse_table()
                continue

            # Checklist item
            if re.match(r'^-\s*\[[ x]\]\s+', line.strip()):
                self._parse_checklist()
                continue

            # Unordered list
            if re.match(r'^[-*]\s+', line.strip()):
                self._parse_list()
                continue

            # Ordered list
            if re.match(r'^\d+\.\s+', line.strip()):
                self._parse_ordered_list()
                continue

            # Bold question pattern (for fixation questions)
            if re.match(r'^\*\*\d+\.', line.strip()):
                self.blocks.append({'type': 'question', 'text': line.strip()})
                self.pos += 1
                continue

            # Regular paragraph
            self._parse_paragraph()

        return self.blocks

    def _parse_code_block(self):
        first_line = self.lines[self.pos].strip()
        lang = first_line[3:].strip()
        self.pos += 1
        code_lines = []
        while self.pos < len(self.lines) and not self.lines[self.pos].strip().startswith('```'):
            code_lines.append(self.lines[self.pos])
            self.pos += 1
        if self.pos < len(self.lines):
            self.pos += 1  # skip closing ```
        self.blocks.append({'type': 'code', 'lang': lang, 'lines': code_lines})

    def _parse_blockquote(self):
        lines = []
        while self.pos < len(self.lines) and self.lines[self.pos].strip().startswith('>'):
            text = re.sub(r'^>\s?', '', self.lines[self.pos])
            lines.append(text)
            self.pos += 1
        full_text = '\n'.join(lines)
        # Determine callout type
        box_type = 'info'
        if re.search(r'\*\*Atenção[:\*]', full_text) or re.search(r'\*\*Importante[:\*]', full_text):
            box_type = 'warning'
        elif re.search(r'\*\*Regra prática[:\*]', full_text) or re.search(r'\*\*Dica[:\*]', full_text):
            box_type = 'tip'
        elif re.search(r'\*\*Nota[:\*]', full_text) or re.search(r'\*\*Nota sobre', full_text):
            box_type = 'info'
        self.blocks.append({'type': 'blockquote', 'text': full_text, 'box_type': box_type})

    def _parse_table(self):
        rows = []
        while self.pos < len(self.lines) and '|' in self.lines[self.pos]:
            cells = [c.strip() for c in self.lines[self.pos].strip().strip('|').split('|')]
            rows.append(cells)
            self.pos += 1
        # Remove separator row (row 1 with dashes)
        if len(rows) > 1 and all(re.match(r'^[\s:]*-+[\s:]*$', c) for c in rows[1]):
            rows.pop(1)
        self.blocks.append({'type': 'table', 'rows': rows})

    def _parse_checklist(self):
        items = []
        while self.pos < len(self.lines) and re.match(r'^-\s*\[[ x]\]\s+', self.lines[self.pos].strip()):
            text = re.sub(r'^-\s*\[[ x]\]\s+', '', self.lines[self.pos].strip())
            items.append(text)
            self.pos += 1
        self.blocks.append({'type': 'checklist', 'items': items})

    def _parse_list(self):
        items = []
        while self.pos < len(self.lines):
            line = self.lines[self.pos]
            m = re.match(r'^(\s*)[-*]\s+(.*)', line)
            if m:
                indent = len(m.group(1))
                items.append({'text': m.group(2), 'indent': indent})
                self.pos += 1
            elif line.strip() and items and not re.match(r'^#{1,6}\s', line) and not line.strip().startswith('```'):
                # Continuation line
                items[-1]['text'] += ' ' + line.strip()
                self.pos += 1
            else:
                break
        self.blocks.append({'type': 'list', 'items': items})

    def _parse_ordered_list(self):
        items = []
        while self.pos < len(self.lines):
            line = self.lines[self.pos]
            m = re.match(r'^\d+\.\s+(.*)', line.strip())
            if m:
                items.append(m.group(1))
                self.pos += 1
            elif line.strip() and items and not re.match(r'^#{1,6}\s', line):
                items[-1] += ' ' + line.strip()
                self.pos += 1
            else:
                break
        self.blocks.append({'type': 'ordered_list', 'items': items})

    def _parse_paragraph(self):
        lines = []
        while self.pos < len(self.lines):
            line = self.lines[self.pos]
            if not line.strip():
                self.pos += 1
                break
            if re.match(r'^#{1,6}\s', line) or line.strip().startswith('```') or line.strip().startswith('>'):
                break
            if '|' in line and self.pos + 1 < len(self.lines) and re.search(r'\|[\s:]*-+', self.lines[self.pos + 1]):
                break
            if re.match(r'^-\s*\[[ x]\]', line.strip()) or re.match(r'^[-*]\s+', line.strip()):
                break
            if re.match(r'^\d+\.\s+', line.strip()):
                break
            if re.match(r'^---+\s*$', line.strip()):
                break
            lines.append(line)
            self.pos += 1
        if lines:
            self.blocks.append({'type': 'paragraph', 'text': '\n'.join(lines)})


# ──────────────────────────────────────────────
# Inline formatting renderer
# ──────────────────────────────────────────────

def render_inline(paragraph, text, base_size=11, base_color=C_BODY, base_bold=False, base_italic=False, font_name='Arial'):
    """Render inline markdown (bold, italic, code, links) into a paragraph."""
    # Pattern for inline elements
    pattern = re.compile(
        r'(\*\*\*(.+?)\*\*\*)'   # bold+italic
        r'|(\*\*(.+?)\*\*)'      # bold
        r'|(\*(.+?)\*)'          # italic
        r'|(`([^`]+)`)'          # inline code
        r'|(\[([^\]]+)\]\(([^)]+)\))'  # link
    )

    pos = 0
    for m in pattern.finditer(text):
        # Add text before match
        if m.start() > pos:
            add_run(paragraph, text[pos:m.start()], bold=base_bold, italic=base_italic,
                    size=base_size, color=base_color, font_name=font_name)

        if m.group(2):  # bold+italic
            add_run(paragraph, m.group(2), bold=True, italic=True,
                    size=base_size, color=base_color, font_name=font_name)
        elif m.group(4):  # bold
            add_run(paragraph, m.group(4), bold=True, italic=base_italic,
                    size=base_size, color=base_color, font_name=font_name)
        elif m.group(6):  # italic
            add_run(paragraph, m.group(6), bold=base_bold, italic=True,
                    size=base_size, color=base_color, font_name=font_name)
        elif m.group(8):  # inline code
            add_run(paragraph, m.group(8), bold=False, italic=False,
                    size=base_size, color=base_color, font_name='Courier New')
        elif m.group(10):  # link
            link_text = m.group(10)
            link_url = m.group(11)
            if link_url.startswith('#'):
                # Internal link
                add_hyperlink(paragraph, link_url[1:], link_text)
            else:
                add_run(paragraph, link_text, bold=base_bold, italic=base_italic,
                        size=base_size, color='0563C1', font_name=font_name)

        pos = m.end()

    # Remaining text
    if pos < len(text):
        add_run(paragraph, text[pos:], bold=base_bold, italic=base_italic,
                size=base_size, color=base_color, font_name=font_name)


def render_cell_inline(cell, text, size=10, bold=False, color=None, font_name='Arial'):
    """Render inline markdown into a table cell's first paragraph."""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    render_inline(p, text, base_size=size, base_color=color or C_BODY, base_bold=bold, font_name=font_name)


# ──────────────────────────────────────────────
# Document Builder
# ──────────────────────────────────────────────

class DocxBeautifier:
    def __init__(self, blocks):
        self.doc = Document()
        self.blocks = blocks
        self.bookmarks = {}  # name -> id
        self._anchor_map = {}  # normalized_pandoc_anchor -> bookmark_name (for TOC resolution)
        self._precompute_anchor_map()
        self._setup_styles()
        self._setup_page()

    def _precompute_anchor_map(self):
        """Pre-scan all H2+ headings and build a map from any possible
        pandoc-style anchor to the sanitized bookmark name we will create.

        Maps: normalized_pandoc_anchor -> sanitized_bookmark_name
        """
        for block in self.blocks:
            if block['type'] != 'heading' or block['level'] < 2:
                continue
            text = block['text']
            level = block['level']
            # The pandoc-style anchor (with accents etc.)
            pandoc = self._generate_pandoc_anchor(text)
            # The actual bookmark name (ASCII, max 40 chars)
            sanitized = _sanitize_bookmark(pandoc) if pandoc else None
            if pandoc and sanitized:
                self._anchor_map[pandoc] = sanitized
            # Short bookmarks (cap3, lab31, apendice_a) only for H2
            if level == 2:
                short = self._generate_bookmark(text)
                if short:
                    self._anchor_map[short] = sanitized or _sanitize_bookmark(short)

    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        # Different first page for header/footer
        section.different_first_page_header_footer = True

    def _setup_styles(self):
        style = self.doc.styles['Normal']
        style.font.name = 'Arial'
        style.font.size = Pt(11)
        style.font.color.rgb = hex_to_rgb(C_BODY)

        # Heading 1
        h1 = self.doc.styles['Heading 1']
        h1.font.name = 'Arial'
        h1.font.size = Pt(18)
        h1.font.bold = True
        h1.font.color.rgb = hex_to_rgb(C_PRIMARY)
        h1.paragraph_format.space_before = Pt(16)
        h1.paragraph_format.space_after = Pt(8)

        # Heading 2
        h2 = self.doc.styles['Heading 2']
        h2.font.name = 'Arial'
        h2.font.size = Pt(13)
        h2.font.bold = True
        h2.font.color.rgb = hex_to_rgb(C_PRIMARY)
        h2.paragraph_format.space_before = Pt(12)
        h2.paragraph_format.space_after = Pt(6)

        # Heading 3
        h3 = self.doc.styles['Heading 3']
        h3.font.name = 'Arial'
        h3.font.size = Pt(12)
        h3.font.bold = False
        h3.font.color.rgb = hex_to_rgb(C_HEADING3)

        # Heading 4
        h4 = self.doc.styles['Heading 4']
        h4.font.name = 'Arial'
        h4.font.size = Pt(11)
        h4.font.bold = False
        h4.font.italic = True
        h4.font.color.rgb = hex_to_rgb(C_ACCENT)

    def build(self):
        i = 0
        title_done = False
        while i < len(self.blocks):
            block = self.blocks[i]

            if block['type'] == 'heading':
                level = block['level']
                text = block['text']

                if level == 1 and not title_done:
                    self._build_title_page(i)
                    title_done = True
                    i += 1
                    continue

                if level == 2:
                    self._add_heading2(text)
                elif level == 3:
                    self._add_heading3(text)
                elif level == 4:
                    self._add_heading4(text)
                elif level == 5:
                    p = self.doc.add_paragraph()
                    render_inline(p, text, base_size=11, base_color=C_ACCENT, base_italic=True)
                else:
                    p = self.doc.add_paragraph()
                    render_inline(p, text, base_size=11, base_bold=True)

            elif block['type'] == 'paragraph':
                p = self.doc.add_paragraph()
                render_inline(p, block['text'])

            elif block['type'] == 'code':
                self._add_code_block(block['lines'])

            elif block['type'] == 'blockquote':
                self._add_callout(block['text'], block['box_type'])

            elif block['type'] == 'table':
                self._add_table(block['rows'])

            elif block['type'] == 'checklist':
                self._add_checklist(block['items'])

            elif block['type'] == 'list':
                self._add_list(block['items'])

            elif block['type'] == 'ordered_list':
                self._add_ordered_list(block['items'])

            elif block['type'] == 'question':
                self._add_question(block['text'])

            elif block['type'] == 'hr':
                self._add_separator()

            i += 1

        self._add_footer()
        return self.doc

    # ── Title page ──

    def _build_title_page(self, start_idx):
        block = self.blocks[start_idx]
        title_text = block['text']

        # Empty space
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(20)

        # Main title
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)
        add_run(p, title_text, bold=True, size=20, color=C_PRIMARY, font_name='Arial')

        # Look ahead for subtitle/metadata
        j = start_idx + 1
        while j < len(self.blocks):
            b = self.blocks[j]
            if b['type'] == 'heading' and b['level'] == 2:
                # Subtitle (e.g., "Segurança Defensiva 2026-1")
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(8)
                add_run(p, b['text'], bold=False, size=14, color=C_SUBTITLE, font_name='Arial')
                j += 1
                continue
            elif b['type'] == 'heading' and b['level'] == 3:
                # Book reference
                text = b['text']
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(12)
                # Handle *text* italic markers
                render_inline(p, text, base_size=12, base_color=C_SUBTITLE, base_italic=True)
                j += 1
                continue
            elif b['type'] == 'blockquote':
                # Version info - render as a note on title page
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(12)
                # Simplified version text
                first_line = b['text'].split('\n')[0].strip()
                render_inline(p, first_line, base_size=10, base_color=C_NOTE, base_italic=True)
                j += 1
                continue
            elif b['type'] == 'hr':
                j += 1
                break
            else:
                break

        # Separator
        p = self.doc.add_paragraph()
        set_paragraph_border_bottom(p, BORDER_CHAPTER)
        p.paragraph_format.space_before = Pt(16)

        # Page break
        self.doc.add_page_break()

        # Build TOC if next block is heading "Sumário"
        if j < len(self.blocks) and self.blocks[j]['type'] == 'heading' and 'Sumário' in self.blocks[j].get('text', ''):
            j += 1  # skip the heading
            self._build_toc(j)
            # Skip TOC blocks (list items until next hr or heading)
            while j < len(self.blocks) and self.blocks[j]['type'] not in ('hr',):
                if self.blocks[j]['type'] == 'heading' and self.blocks[j]['level'] <= 2 and 'Sumário' not in self.blocks[j].get('text', ''):
                    break
                j += 1
            if j < len(self.blocks) and self.blocks[j]['type'] == 'hr':
                j += 1  # skip hr after toc

        # Update position - we need to adjust. Since we return to the main loop,
        # we advance via mutating self.blocks by marking them processed.
        # Instead, we'll just set a flag to skip these blocks
        for k in range(start_idx + 1, min(j, len(self.blocks))):
            self.blocks[k] = {'type': '_skip'}

    def _resolve_toc_anchor(self, raw_anchor):
        """Resolve a TOC anchor from the markdown to a sanitized bookmark name.

        Handles pandoc-style inconsistencies: leading dashes from emojis,
        double dashes from special chars, etc.
        Returns an ASCII-safe, Word-compatible bookmark name.
        """
        # Normalize: strip leading/trailing dashes, collapse double dashes
        normalized = re.sub(r'-+', '-', raw_anchor).strip('-')
        # Direct match in anchor map -> returns sanitized name
        if normalized in self._anchor_map:
            return self._anchor_map[normalized]
        # Try substring matching against known anchors
        for known_anchor in self._anchor_map:
            if normalized in known_anchor or known_anchor in normalized:
                return self._anchor_map[known_anchor]
        # Fallback: sanitize the normalized anchor directly
        log.warning("TOC anchor not resolved: %s (normalized: %s)", raw_anchor, normalized)
        return _sanitize_bookmark(normalized)

    def _build_toc(self, start_idx):
        """Build table of contents from list blocks."""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        add_run(p, 'Sumário', bold=True, size=16, color=C_PRIMARY, font_name='Arial')

        # Collect TOC entries from list blocks
        toc_entries = []
        j = start_idx
        while j < len(self.blocks):
            b = self.blocks[j]
            if b['type'] == 'list':
                for item in b['items']:
                    toc_entries.append(item)
                j += 1
            elif b['type'] == 'hr':
                break
            elif b['type'] == 'heading':
                break
            else:
                j += 1

        # Render TOC entries with background
        for entry in toc_entries:
            text = entry['text']
            indent = entry.get('indent', 0)
            p = self.doc.add_paragraph()
            set_paragraph_bg(p, BG_HEADER_LIGHT)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if indent > 0:
                p.paragraph_format.left_indent = Cm(1.0)

            # Extract link if present
            link_match = re.match(r'\[(.+?)\]\(#(.+?)\)', text)
            if link_match:
                link_text = link_match.group(1)
                raw_anchor = link_match.group(2)
                resolved = self._resolve_toc_anchor(raw_anchor)
                add_hyperlink(p, resolved, link_text)
            else:
                render_inline(p, text, base_size=11, base_color=C_PRIMARY)

        self.doc.add_page_break()

    # ── Headings ──

    def _add_heading2(self, text):
        """## headings → Heading1 with bookmark."""
        # Detect lab heading with emoji
        emoji_level = None
        clean_text = text
        for emoji, level_name in [('🟢', 'Iniciante'), ('🟡', 'Intermediário'), ('🔴', 'Avançado')]:
            if emoji in text:
                emoji_level = level_name
                clean_text = text.replace(emoji, '').strip()
                break

        # Build paragraph manually to control run order (badge before text)
        p = self.doc.add_paragraph(style='Heading 1')
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)

        # Add badge FIRST if it's a lab heading
        if emoji_level:
            add_badge(p, emoji_level)
            add_run(p, '  ', size=18, font_name='Arial')

        # Add heading text
        add_run(p, clean_text, bold=True, size=18, color=C_PRIMARY, font_name='Arial')

        # Add sanitized bookmarks — both short name AND pandoc-style anchor for TOC links
        bookmark_name = self._generate_bookmark(text)
        if bookmark_name:
            safe_name = _sanitize_bookmark(bookmark_name)
            bid = next_bookmark_id()
            add_bookmark(p, safe_name, bid)
            self.bookmarks[safe_name] = bid

        # Also add pandoc-style anchor (sanitized) so TOC hyperlinks work
        pandoc_anchor = self._generate_pandoc_anchor(text)
        if pandoc_anchor:
            safe_pandoc = _sanitize_bookmark(pandoc_anchor)
            if safe_pandoc and safe_pandoc not in self.bookmarks:
                bid2 = next_bookmark_id()
                add_bookmark(p, safe_pandoc, bid2)
                self.bookmarks[safe_pandoc] = bid2

        # Chapter separator
        set_paragraph_border_bottom(p, BORDER_CHAPTER)

    def _add_heading3(self, text):
        """### headings → Heading2."""
        p = self.doc.add_paragraph(style='Heading 2')
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        add_run(p, text, bold=True, size=13, color=C_PRIMARY, font_name='Arial')
        # Add sanitized pandoc-style bookmark
        pandoc_anchor = self._generate_pandoc_anchor(text)
        if pandoc_anchor:
            safe = _sanitize_bookmark(pandoc_anchor)
            if safe and safe not in self.bookmarks:
                bid = next_bookmark_id()
                add_bookmark(p, safe, bid)
                self.bookmarks[safe] = bid

    def _add_heading4(self, text):
        """#### headings → Heading3."""
        p = self.doc.add_paragraph(style='Heading 3')
        add_run(p, text, size=12, color=C_HEADING3, font_name='Arial')

    def _generate_bookmark(self, text):
        """Generate short bookmark name from heading text."""
        # Chapter heading
        m = re.match(r'.*Capítulo\s+(\d+)', text)
        if m:
            return f'cap{m.group(1)}'
        # Lab heading
        m = re.search(r'LAB\s+(\d+)\.(\d+)', text)
        if m:
            return f'lab{m.group(1)}{m.group(2)}'
        # Appendix
        if 'Apêndice A' in text or 'Apêndice A' in text:
            return 'apendice_a'
        return None

    @staticmethod
    def _generate_pandoc_anchor(text):
        """Generate pandoc-style anchor from heading text (for TOC link compatibility)."""
        # Strip emojis
        anchor = text
        for emoji in ('🟢', '🟡', '🔴', '✔'):
            anchor = anchor.replace(emoji, '')
        anchor = anchor.lower().strip()
        # Remove all punctuation except hyphens, underscores, spaces
        anchor = re.sub(r'[^\w\s-]', '', anchor)
        # Replace whitespace with hyphens
        anchor = re.sub(r'\s+', '-', anchor)
        # Collapse multiple hyphens
        anchor = re.sub(r'-+', '-', anchor)
        # Strip leading/trailing hyphens
        anchor = anchor.strip('-')
        return anchor

    # ── Code blocks ──

    def _add_code_block(self, lines):
        for line in lines:
            p = self.doc.add_paragraph()
            set_paragraph_bg(p, BG_CODE)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(1.0)

            # Color bash comments
            if line.strip().startswith('#'):
                add_run(p, line, size=9, color=C_MUTED, font_name='Courier New')
            else:
                add_run(p, line if line else ' ', size=9, color=C_BODY, font_name='Courier New')

    # ── Callout boxes ──

    def _add_callout(self, text, box_type):
        bg_map = {
            'info': BG_INFO,
            'tip': BG_TIP,
            'warning': BG_WARNING,
            'danger': BG_DANGER,
        }
        bg = bg_map.get(box_type, BG_INFO)

        # Split into lines for multi-line callouts
        for line_text in text.split('\n'):
            p = self.doc.add_paragraph()
            set_paragraph_bg(p, bg)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Cm(0.5)

            if box_type == 'tip':
                set_paragraph_border_left(p, BORDER_TIP)

            render_inline(p, line_text.strip(), base_size=10, base_color=C_NOTE)

    # ── Tables ──

    def _add_table(self, rows):
        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        # Pad rows
        for r in rows:
            while len(r) < num_cols:
                r.append('')

        # Detect special table types
        is_troubleshooting = any('Problema' in c and 'Solução' in c for r in rows[:1] for c in r if isinstance(c, str))
        is_criteria = any('Critério' in c and 'Pontos' in c for r in rows[:1] for c in r if isinstance(c, str))

        table = self.doc.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_table_borders(table)

        # Header row
        for ci, cell_text in enumerate(rows[0]):
            cell = table.cell(0, ci)
            if is_troubleshooting:
                set_cell_bg(cell, BG_DANGER)
            elif is_criteria:
                set_cell_bg(cell, BG_HEADER_LIGHT)
            else:
                set_cell_bg(cell, BG_HEADER_LIGHT)
            render_cell_inline(cell, cell_text, size=10, bold=True)

        # Data rows
        for ri in range(1, len(rows)):
            is_total = any('TOTAL' in c.upper() for c in rows[ri])
            for ci, cell_text in enumerate(rows[ri]):
                cell = table.cell(ri, ci)
                if is_total:
                    set_cell_bg(cell, BG_HEADER_LIGHT)
                    render_cell_inline(cell, cell_text, size=10, bold=True)
                elif ri % 2 == 0 and len(rows) > 5:
                    set_cell_bg(cell, BG_TABLE_ALT)
                    render_cell_inline(cell, cell_text, size=10)
                else:
                    render_cell_inline(cell, cell_text, size=10)

    # ── Checklist ──

    def _add_checklist(self, items):
        for item in items:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_run(p, '☐  ', size=10, font_name='Arial')

            # Check for (check `xxx`) pattern
            check_match = re.search(r'\(check\s+`([^`]+)`\)', item)
            if check_match:
                before = item[:check_match.start()].strip()
                check_id = check_match.group(1)
                render_inline(p, before, base_size=10)
                add_run(p, f'  (check ', size=9, color=C_MUTED)
                add_run(p, check_id, size=9, color=C_BODY, font_name='Courier New')
                add_run(p, ')', size=9, color=C_MUTED)
            else:
                render_inline(p, item, base_size=10)

    # ── Lists ──

    def _add_list(self, items):
        for item in items:
            p = self.doc.add_paragraph()
            indent_level = item.get('indent', 0) // 2
            p.paragraph_format.left_indent = Cm(0.5 + indent_level * 0.5)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_run(p, '•  ', size=11, font_name='Arial')
            render_inline(p, item['text'])

    def _add_ordered_list(self, items):
        for idx, text in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_run(p, f'{idx}. ', bold=True, size=11, font_name='Arial')
            render_inline(p, text)

    # ── Questions ──

    def _add_question(self, text):
        """Add fixation question with answer box."""
        # Strip ** markers
        clean = re.sub(r'\*\*', '', text)

        # Check for (Desafio) prefix
        is_challenge = '(Desafio)' in clean

        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)

        if is_challenge:
            # Split at (Desafio)
            parts = clean.split('(Desafio)')
            render_inline(p, parts[0].strip(), base_size=10, base_bold=True)
            add_run(p, ' (Desafio)', bold=True, size=10, color='C45911')
            if len(parts) > 1:
                render_inline(p, parts[1].strip(), base_size=10, base_bold=True)
        else:
            render_inline(p, clean, base_size=10, base_bold=True)

        # Answer box (1x1 table)
        table = self.doc.add_table(rows=1, cols=1)
        add_table_borders(table, color='AAAAAA')
        cell = table.cell(0, 0)
        set_cell_bg(cell, BG_ANSWER)
        p = cell.paragraphs[0]
        add_run(p, 'Resposta (até 4 linhas):', italic=True, size=9, color=C_PLACEHOLDER, font_name='Arial')
        for _ in range(3):
            cell.add_paragraph()

    # ── Separators ──

    def _add_separator(self):
        p = self.doc.add_paragraph()
        set_paragraph_border_bottom(p, BORDER_SECTION)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)

    # ── Footer ──

    def _add_footer(self):
        section = self.doc.sections[0]
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, 'Segurança Defensiva 2026-1 — Laboratórios', size=9, color=C_MUTED, font_name='Arial')

        # Add page number field
        run = p.add_run('    ')
        fld_xml = (
            f'<w:fldSimple {nsdecls("w")} w:instr=" PAGE \\* MERGEFORMAT ">'
            f'  <w:r><w:rPr><w:sz w:val="18"/><w:color w:val="{C_MUTED}"/></w:rPr><w:t>1</w:t></w:r>'
            f'</w:fldSimple>'
        )
        p._p.append(parse_xml(fld_xml))


# ──────────────────────────────────────────────
# Main
# ──────────────────────────────────────────────

def convert(input_path: str, output_path: str = None) -> str:
    input_file = Path(input_path).resolve()
    if not input_file.exists():
        raise FileNotFoundError(f"Input file not found: {input_file}")
    if input_file.suffix.lower() not in ('.md', '.markdown'):
        raise ValueError(f"Input must be a markdown file: {input_file}")

    if output_path:
        output_file = Path(output_path).resolve()
    else:
        output_file = input_file.with_suffix('.docx')

    log.info("Converting %s -> %s", input_file, output_file)

    content = input_file.read_text(encoding='utf-8')
    parser = MarkdownParser(content)
    blocks = parser.parse()

    builder = DocxBeautifier(blocks)
    doc = builder.build()
    doc.save(str(output_file))

    log.info("Done: %s", output_file)
    return str(output_file)


def main():
    parser = argparse.ArgumentParser(
        description='Convert markdown to beautifully formatted DOCX.'
    )
    parser.add_argument('input_file', help='Path to the input markdown file')
    parser.add_argument('-o', '--output', help='Output DOCX file path (default: same name as input with .docx)')
    args = parser.parse_args()

    try:
        output = convert(args.input_file, args.output)
        print(f"Output: {output}")
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
